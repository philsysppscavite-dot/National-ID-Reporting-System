from __future__ import annotations

import io
import json
import sqlite3
from datetime import datetime
from pathlib import Path

from flask import abort, flash, redirect, render_template, request, send_file, send_from_directory, session, url_for
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from werkzeug.utils import secure_filename

from .auth import current_role, current_user_id, is_isa_or_admin, roles_required
from .db import ROLE_ADMIN, ROLE_ISA, ROLE_USER, ROLES
from .repository import (
    CITY_MUNICIPALITIES,
    SERVICE_TYPES,
    NID_CONCERN_TYPES,
    NID_CONCERN_STATUSES,
    GENDER_OPTIONS,
    DIGITAL_ID_ASSISTANCE_OPTIONS,
    YES_NO_OPTIONS,
    OVERSEAS_REGISTRANT_OPTIONS,
    AUTHENTICATED_OPTIONS,
    DOB_MONTH_OPTIONS,
    add_concern_message,
    create_nid_concern,
    create_user,
    delete_data_entry,
    delete_schedule,
    output_summary_grand_totals,
    delete_employee,
    delete_output,
    delete_signatory,
    delete_nid_concern,
    employee_output_totals,
    fetch_settings,
    get_data_entry,
    get_employee,
    get_entry_defaults,
    get_output,
    get_signatory,
    get_nid_concern,
    get_user,
    city_service_rows,
    delete_imported_outputs,
    delete_import_source,
    delete_outputs_by_source_ref,
    set_import_source_active,
    get_import_source,
    get_schedule,
    is_valid_trn_ref_no,
    list_active_users,
    list_concern_messages,
    list_conversation,
    list_data_entries,
    list_employees,
    list_import_sources,
    list_nid_concerns,
    list_outputs,
    list_pending_client_notifications,
    list_schedule_employees,
    list_schedules,
    list_signatories,
    list_users,
    mark_concern_client_informed,
    mark_data_entries_sent_to_sheet,
    monthly_city_records,
    NID_CONCERN_STATUS_PROGRESS,
    nid_concern_matrix_summary,
    nid_concern_monthly_matrix,
    nid_concern_monthly_trend,
    nid_concern_progress,
    output_per_person_rows,
    resolve_entry_defaults_for_form,
    save_data_entry,
    save_employee,
    save_entry_defaults,
    save_output,
    save_schedule,
    save_signatory,
    send_direct_message,
    set_setting,
    signatories_for_report,
    unread_dm_counts_by_sender,
    update_nid_concern,
    update_nid_concern_status,
    update_user,
    user_can_view_concern,
)
from .services.export import ExportError, export_data_entries_to_sheet
from .services.importers import import_from_apps_script, import_from_csv_url, reimport_source
from .services.reports import (
    build_dar_workbook,
    generate_all_employees_dar_pdf_zip,
    generate_bulk_reports_by_date,
    generate_city_service_summary_report,
    generate_employee_report_by_date,
    generate_locator_chart_all_report,
    generate_locator_chart_report,
    generate_output_summary_report,
    generate_all_employees_output_reports_by_date,
    generate_schedule_report,
)

def register_routes(app):
    @app.get("/health")
    def health_check():
        return {"status": "ok"}, 200

    @app.context_processor
    def inject_settings():
        current_uid = session.get("user_id")
        unread_total = 0
        pending_client_notifications_total = 0
        if current_uid:
            unread_total = sum(unread_dm_counts_by_sender(current_uid).values())
            pending_client_notifications_total = len(list_pending_client_notifications(current_uid))
        return {
            "app_settings": fetch_settings(),
            "city_municipalities": CITY_MUNICIPALITIES,
            "ROLE_ADMIN": ROLE_ADMIN,
            "ROLE_ISA": ROLE_ISA,
            "current_role": session.get("role"),
            "current_full_name": session.get("full_name"),
            "unread_dm_total": unread_total,
            "pending_client_notifications_total": pending_client_notifications_total,
        }

    @app.get("/logo/<path:filename>")
    def logo_file(filename):
        # Try uploaded logo first (uploads stored in UPLOAD_DIR), then fall back to the repo /logo directory
        upload_dir = Path(app.config.get("UPLOAD_DIR", Path(app.root_path).parent))
        uploaded_candidate = upload_dir / filename
        if uploaded_candidate.exists():
            return send_from_directory(str(upload_dir), filename)

        fixed_dir = Path(app.root_path).parent / "logo"
        fixed_candidate = fixed_dir / filename
        if fixed_candidate.exists():
            return send_from_directory(str(fixed_dir), filename)

        # If an uploaded logo exists under a generic name, serve it as a fallback
        for candidate in upload_dir.glob("logo.*"):
            return send_from_directory(str(upload_dir), candidate.name)

        from flask import abort

        abort(404)

    @app.route("/")
    def dashboard():
        try:
            start_date, end_date = _resolve_date_filter()
            outputs = list_outputs({"start_date": start_date, "end_date": end_date})
            employees = list_employees()
            total_quantity = sum(row["quantity"] for row in outputs)
            active_employees = sum(1 for row in employees if row["active"])
            city_rows = monthly_city_records(start_date[:7])

            # Recent Tickets filters (independent of the main output date range above)
            ticket_date_from = (request.args.get("ticket_date_from") or "").strip()
            ticket_date_to = (request.args.get("ticket_date_to") or "").strip()
            ticket_type = (request.args.get("ticket_type") or "").strip()
            ticket_status = (request.args.get("ticket_status") or "").strip()
            ticket_search = (request.args.get("ticket_search") or "").strip()
            ticket_filters_active = bool(
                ticket_date_from or ticket_date_to or ticket_type or ticket_status or ticket_search
            )

            concern_filters: dict = {}
            if ticket_date_from:
                concern_filters["start_date"] = ticket_date_from
            if ticket_date_to:
                concern_filters["end_date"] = ticket_date_to
            if ticket_type:
                concern_filters["concern_type"] = ticket_type

            all_concerns = list_nid_concerns(concern_filters)
            if current_role() not in (ROLE_ADMIN, ROLE_ISA):
                all_concerns = [c for c in all_concerns if c["reported_by_user_id"] == current_user_id()]

            if ticket_status:
                all_concerns = [c for c in all_concerns if c["status"] == ticket_status]

            if ticket_search:
                needle = ticket_search.lower()
                def _matches(c):
                    haystack = " ".join(
                        str(c.get(field) or "")
                        for field in ("trn_or_ref_no", "description", "reported_by", "concern_type", "status")
                    ).lower()
                    return needle in haystack
                all_concerns = [c for c in all_concerns if _matches(c)]

            # Pagination -- every ticket is reachable, just paged instead of
            # silently truncated to a "recent 8".
            TICKETS_PER_PAGE = 10
            total_tickets = len(all_concerns)
            total_pages = max(1, -(-total_tickets // TICKETS_PER_PAGE))  # ceil div
            try:
                ticket_page = int(request.args.get("ticket_page", 1))
            except ValueError:
                ticket_page = 1
            ticket_page = min(max(ticket_page, 1), total_pages)
            page_start = (ticket_page - 1) * TICKETS_PER_PAGE
            page_end = page_start + TICKETS_PER_PAGE
            recent_nid_concerns = all_concerns[page_start:page_end]

            # If the last "Log a New Ticket" submit failed validation, this
            # carries the user's typed values back so the form can be
            # repopulated instead of coming up blank. Pop it so it's only
            # used once.
            ticket_form_old = session.pop("ticket_form_old", None)

            return render_template(
                "dashboard.html",
                month=start_date[:7],
                start_date=start_date,
                end_date=end_date,
                outputs=outputs,
                active_employees=active_employees,
                total_quantity=total_quantity,
                city_rows=city_rows,
                import_sources=list_import_sources(),
                nid_concern_matrix=nid_concern_matrix_summary(),
                nid_concern_statuses=NID_CONCERN_STATUSES,
                nid_concern_types=NID_CONCERN_TYPES,
                nid_concern_progress_map=NID_CONCERN_STATUS_PROGRESS,
                recent_nid_concerns=recent_nid_concerns,
                pending_client_notifications=list_pending_client_notifications(current_user_id()),
                ticket_date_from=ticket_date_from,
                ticket_date_to=ticket_date_to,
                ticket_type=ticket_type,
                ticket_status=ticket_status,
                ticket_search=ticket_search,
                ticket_filters_active=ticket_filters_active,
                ticket_page=ticket_page,
                total_pages=total_pages,
                total_tickets=total_tickets,
                ticket_page_start=(page_start + 1) if total_tickets else 0,
                ticket_page_end=min(page_end, total_tickets),
                ticket_form_old=ticket_form_old,
            )
        except Exception as e:
            from flask import jsonify
            return jsonify({"error": str(e)}), 500

    @app.post("/nid-concerns/new")
    def nid_concern_create():
        try:
            new_id = create_nid_concern(
                request.form,
                reporter_user_id=current_user_id(),
                reporter_full_name=session.get("full_name"),
            )
            flash("Ticket logged.", "success")
            return redirect(url_for("nid_concern_detail", concern_id=new_id))
        except ValueError as exc:
            flash(str(exc), "error")
            # Keep whatever the user already typed so a validation error
            # (bad TRN length, bad mobile number, etc.) doesn't force them
            # to re-type the whole ticket from scratch on the next load.
            session["ticket_form_old"] = {
                "date_reported": request.form.get("date_reported", ""),
                "trn_or_ref_no": request.form.get("trn_or_ref_no", ""),
                "concern_type": request.form.get("concern_type", ""),
                "mobile_number": request.form.get("mobile_number", ""),
                "description": request.form.get("description", ""),
            }
            return redirect(url_for("dashboard"))

    @app.post("/nid-concerns/<int:concern_id>/inform-client")
    def nid_concern_inform_client(concern_id):
        concern = get_nid_concern(concern_id)
        if not concern:
            abort(404)
        if not user_can_view_concern(concern, current_user_id(), current_role()):
            abort(403)
        if concern["status"] != "Resolved":
            flash("This ticket isn't marked Resolved yet.", "error")
        else:
            mark_concern_client_informed(concern_id)
            flash("Thanks -- marked as informed.", "success")
        next_url = request.form.get("next") or url_for("dashboard")
        return redirect(next_url)

    @app.route("/nid-concerns/<int:concern_id>/edit", methods=["GET", "POST"])
    def nid_concern_edit(concern_id):
        concern = get_nid_concern(concern_id)
        if not concern:
            abort(404)
        if not user_can_view_concern(concern, current_user_id(), current_role()):
            abort(403)
        if request.method == "POST":
            try:
                update_nid_concern(concern_id, request.form)
                flash("Ticket updated.", "success")
                return redirect(url_for("dashboard"))
            except ValueError as exc:
                flash(str(exc), "error")
                concern = dict(concern)
                concern.update(request.form)
        return render_template(
            "ticket_form.html",
            concern=concern,
            nid_concern_types=NID_CONCERN_TYPES,
        )

    @app.post("/nid-concerns/<int:concern_id>/status")
    @roles_required(ROLE_ISA)
    def nid_concern_update_status(concern_id):
        concern = get_nid_concern(concern_id)
        if concern:
            new_status = request.form.get("status") or concern["status"]
            if new_status not in NID_CONCERN_STATUSES:
                flash("Unknown status.", "error")
            else:
                update_nid_concern_status(concern_id, new_status)
                flash("Ticket status updated.", "success")
        next_url = request.form.get("next") or url_for("dashboard")
        return redirect(next_url)

    @app.post("/nid-concerns/<int:concern_id>/delete")
    @roles_required(ROLE_ISA)
    def nid_concern_delete(concern_id):
        delete_nid_concern(concern_id)
        flash("Ticket removed.", "success")
        return redirect(url_for("dashboard"))

    @app.route("/nid-concerns/<int:concern_id>", methods=["GET", "POST"])
    def nid_concern_detail(concern_id):
        concern = get_nid_concern(concern_id)
        if not concern:
            abort(404)
        if not user_can_view_concern(concern, current_user_id(), current_role()):
            abort(403)
        if request.method == "POST":
            add_concern_message(concern_id, current_user_id(), request.form.get("body"))
            return redirect(url_for("nid_concern_detail", concern_id=concern_id))
        return render_template(
            "ticket_detail.html",
            concern=concern,
            messages=list_concern_messages(concern_id),
            nid_concern_statuses=NID_CONCERN_STATUSES,
            progress=nid_concern_progress(concern["status"]),
            can_manage=current_role() in (ROLE_ADMIN, ROLE_ISA),
        )

    @app.route("/nid-concerns/report")
    def nid_concern_report():
        month = (request.args.get("month") or datetime.now().strftime("%Y-%m")).strip()
        try:
            trend_months = int(request.args.get("trend_months", 6))
        except ValueError:
            trend_months = 6
        trend_months = min(max(trend_months, 3), 24)

        monthly_matrix = nid_concern_monthly_matrix(month)
        trend = nid_concern_monthly_trend(months=trend_months, end_month=month)
        current_month_stats = next((row for row in trend if row["month"] == month), None)

        return render_template(
            "nid_concern_report.html",
            month=month,
            monthly_matrix=monthly_matrix,
            nid_concern_statuses=NID_CONCERN_STATUSES,
            trend=trend,
            trend_months=trend_months,
            current_month_stats=current_month_stats,
        )

    @app.get("/nid-concerns/report/download")
    def download_nid_concern_report():
        month = (request.args.get("month") or datetime.now().strftime("%Y-%m")).strip()
        try:
            trend_months = int(request.args.get("trend_months", 6))
        except ValueError:
            trend_months = 6
        trend_months = min(max(trend_months, 3), 24)

        monthly_matrix = nid_concern_monthly_matrix(month)
        trend = nid_concern_monthly_trend(months=trend_months, end_month=month)

        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")

        wb = Workbook()

        # Sheet 1: this month's concern-type x status matrix
        ws1 = wb.active
        ws1.title = f"{month} Matrix"[:31]
        headers = ["Concern Type"] + NID_CONCERN_STATUSES + ["Total"]
        for col_idx, header in enumerate(headers, 1):
            cell = ws1.cell(row=1, column=col_idx, value=header)
            cell.font = header_font
            cell.fill = header_fill
        for row_idx, row in enumerate(monthly_matrix, 2):
            ws1.cell(row=row_idx, column=1, value=row["concern_type"])
            for col_idx, status in enumerate(NID_CONCERN_STATUSES, 2):
                ws1.cell(row=row_idx, column=col_idx, value=row.get(status, 0))
            ws1.cell(row=row_idx, column=len(NID_CONCERN_STATUSES) + 2, value=row["total"])
        ws1.column_dimensions["A"].width = 34
        for col_idx in range(2, len(NID_CONCERN_STATUSES) + 3):
            ws1.column_dimensions[chr(64 + col_idx)].width = 14

        # Sheet 2: trend across months
        ws2 = wb.create_sheet(title="Monthly Trend")
        trend_headers = ["Month"] + NID_CONCERN_STATUSES + ["Total", "Resolution Rate (%)"]
        for col_idx, header in enumerate(trend_headers, 1):
            cell = ws2.cell(row=1, column=col_idx, value=header)
            cell.font = header_font
            cell.fill = header_fill
        for row_idx, row in enumerate(trend, 2):
            ws2.cell(row=row_idx, column=1, value=row["month"])
            for col_idx, status in enumerate(NID_CONCERN_STATUSES, 2):
                ws2.cell(row=row_idx, column=col_idx, value=row.get(status, 0))
            ws2.cell(row=row_idx, column=len(NID_CONCERN_STATUSES) + 2, value=row["total"])
            ws2.cell(row=row_idx, column=len(NID_CONCERN_STATUSES) + 3, value=row["resolution_rate"])
        ws2.column_dimensions["A"].width = 14
        for col_idx in range(2, len(NID_CONCERN_STATUSES) + 4):
            ws2.column_dimensions[chr(64 + col_idx)].width = 16

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        filename = f"nid_concerns_report_{month}.xlsx"
        return send_file(
            output,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            as_attachment=True,
            download_name=filename,
        )

    @app.route("/employees")
    def employees():
        return render_template("employees.html", employees=list_employees())

    @app.route("/employees/new", methods=["GET", "POST"])
    @app.route("/employees/<int:employee_id>/edit", methods=["GET", "POST"])
    def employee_form(employee_id=None):
        employee = get_employee(employee_id) if employee_id else None
        if request.method == "POST":
            try:
                save_employee(employee_id, request.form)
                flash("Employee saved successfully.", "success")
                return redirect(url_for("employees"))
            except sqlite3.IntegrityError:
                flash("Employee code already exists. Please use a unique code.", "error")
                employee = request.form
        return render_template("employee_form.html", employee=employee)

    @app.post("/employees/<int:employee_id>/delete")
    def employee_delete(employee_id):
        delete_employee(employee_id)
        flash("Employee deleted.", "success")
        return redirect(url_for("employees"))

    @app.route("/outputs")
    def outputs():
        filters = {
            "employee_id": request.args.get("employee_id") or None,
            "start_date": request.args.get("start_date") or "",
            "end_date": request.args.get("end_date") or "",
        }
        filters = {key: value for key, value in filters.items() if value}
        preview_query = {}
        if filters.get("employee_id"):
            preview_query["employee_id"] = filters["employee_id"]
        if filters.get("start_date"):
            preview_query["start_date"] = filters["start_date"]
        if filters.get("end_date"):
            preview_query["end_date"] = filters["end_date"]
        export_query = {k: filters[k] for k in ("start_date", "end_date") if filters.get(k)}
        outputs = output_per_person_rows(filters)
        total_quantity = sum(row.get("total", 0) for row in outputs)
        return render_template(
            "outputs.html",
            outputs=outputs,
            employees=list_employees(),
            filters=filters,
            preview_query=preview_query,
            export_query=export_query,
            service_types=SERVICE_TYPES,
            total_quantity=total_quantity,
        )

    @app.get("/reports/export-all-outputs")
    def export_all_outputs():
        start_date, end_date = _resolve_output_report_date_range()
        zip_path = generate_all_employees_output_reports_by_date(start_date, end_date)
        return send_file(
            zip_path,
            as_attachment=True,
            download_name=f"{start_date}_to_{end_date}_all_employees_output_reports.zip",
        )

    @app.route("/outputs/new", methods=["GET", "POST"])
    @app.route("/outputs/<int:output_id>/edit", methods=["GET", "POST"])
    def output_form(output_id=None):
        output = get_output(output_id) if output_id else None
        if request.method == "POST":
            save_output(output_id, request.form)
            flash("Employee output saved successfully.", "success")
            return redirect(url_for("outputs"))
        return render_template(
            "output_form.html",
            output=output,
            employees=list_employees(),
        )

    @app.post("/outputs/<int:output_id>/delete")
    def output_delete(output_id):
        delete_output(output_id)
        flash("Employee output deleted.", "success")
        return redirect(url_for("outputs"))

    @app.route("/data-entries")
    def data_entries():
        filters = {
            "city_municipality": request.args.get("city_municipality") or None,
            "rko_employee_id": request.args.get("rko_employee_id") or None,
            "start_date": request.args.get("start_date") or "",
            "end_date": request.args.get("end_date") or "",
            "search": request.args.get("search") or "",
        }
        filters = {key: value for key, value in filters.items() if value}
        entries = list_data_entries(filters)
        return render_template(
            "data_entries.html",
            entries=entries,
            employees=list_employees(),
            filters=filters,
        )

    @app.route("/data-entries/new", methods=["GET", "POST"])
    @app.route("/data-entries/<int:entry_id>/edit", methods=["GET", "POST"])
    def data_entry_form(entry_id=None):
        entry = get_data_entry(entry_id) if entry_id else None
        # Auto-lock only applies to *new* entries -- editing an existing
        # entry always shows/edits its own already-saved values.
        default_values, locked_fields = (
            resolve_entry_defaults_for_form(current_user_id()) if entry_id is None else ({}, set())
        )
        if request.method == "POST":
            try:
                form_data = request.form.to_dict()
                # A locked field is only ever a convenience for the user who
                # locked it -- enforce the saved default server-side so a
                # disabled/read-only field can never be overridden by a
                # tampered submission.
                for field in locked_fields:
                    form_data[field] = default_values.get(field, "")
                save_data_entry(
                    entry_id,
                    form_data,
                    created_by_user_id=current_user_id(),
                    created_by_name=session.get("full_name"),
                )
                flash("Data entry saved successfully.", "success")
                return redirect(url_for("data_entries"))
            except ValueError as exc:
                flash(str(exc), "error")
                entry = form_data
        return render_template(
            "data_entry_form.html",
            entry=entry,
            default_values=default_values,
            locked_fields=locked_fields,
            employees=list_employees(),
            gender_options=GENDER_OPTIONS,
            digital_id_assistance_options=DIGITAL_ID_ASSISTANCE_OPTIONS,
            yes_no_options=YES_NO_OPTIONS,
            overseas_registrant_options=OVERSEAS_REGISTRANT_OPTIONS,
            authenticated_options=AUTHENTICATED_OPTIONS,
            dob_month_options=DOB_MONTH_OPTIONS,
            service_types=SERVICE_TYPES,
        )

    @app.post("/data-entries/<int:entry_id>/delete")
    def data_entry_delete(entry_id):
        delete_data_entry(entry_id)
        flash("Data entry deleted.", "success")
        return redirect(url_for("data_entries"))

    @app.route("/my-entry-defaults", methods=["GET", "POST"])
    def entry_defaults_form():
        user_id = current_user_id()
        if request.method == "POST":
            try:
                save_entry_defaults(user_id, request.form)
                flash("Your entry defaults were saved.", "success")
                return redirect(url_for("entry_defaults_form"))
            except ValueError as exc:
                flash(str(exc), "error")
        return render_template(
            "entry_defaults_form.html",
            defaults=get_entry_defaults(user_id),
            employees=list_employees(),
        )

    @app.post("/data-entries/send-to-sheet")
    def data_entries_send_to_sheet():
        # Regular users send only their own not-yet-sent entries;
        # Administrators/ISA can send everyone's, matching who's allowed to
        # see the full log on the Data Entry page.
        filters = {"sent_to_sheet": False}
        if not is_isa_or_admin():
            filters["created_by_user_id"] = current_user_id()
        entries = list_data_entries(filters)
        if not entries:
            flash("There are no new data entries to send.", "success")
            return redirect(url_for("data_entries"))
        try:
            sheet_url = fetch_settings().get("trn_logsheet_url", "")
            sent_count = export_data_entries_to_sheet(entries, sheet_url)
            mark_data_entries_sent_to_sheet([e["id"] for e in entries])
            flash(f"Sent {sent_count} data entr{'y' if sent_count == 1 else 'ies'} to the Google Sheet.", "success")
        except ExportError as exc:
            flash(str(exc), "error")
        return redirect(url_for("data_entries"))

    @app.route("/output-totals")
    def output_totals():
        filters = {
            "employee_id": request.args.get("employee_id") or None,
            "start_date": request.args.get("start_date") or "",
            "end_date": request.args.get("end_date") or "",
        }
        filters = {key: value for key, value in filters.items() if value}
        rows = employee_output_totals(filters)
        total_quantity = sum(row.get("total", 0) for row in rows)
        return render_template(
            "output_totals.html",
            rows=rows,
            employees=list_employees(),
            filters=filters,
            service_types=SERVICE_TYPES,
            total_quantity=total_quantity,
        )

    @app.post("/output-totals/download")
    def download_output_totals():
        filters = {
            "employee_id": request.form.get("employee_id") or None,
            "start_date": request.form.get("start_date") or "",
            "end_date": request.form.get("end_date") or "",
        }
        filters = {key: value for key, value in filters.items() if value}
        rows = employee_output_totals(filters)
        total_quantity = sum(row.get("total", 0) for row in rows)

        wb = Workbook()
        ws = wb.active
        ws.title = "Output Totals"

        # Header row
        headers = ["Employee"] + SERVICE_TYPES + ["Total"]
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.value = header
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")

        # Data rows
        for row_idx, row in enumerate(rows, 2):
            ws.cell(row=row_idx, column=1).value = row["name"]
            for col_idx, service in enumerate(SERVICE_TYPES, 2):
                ws.cell(row=row_idx, column=col_idx).value = row.get(service, 0)
            ws.cell(row=row_idx, column=len(SERVICE_TYPES) + 2).value = row["total"]

        # Grand total footer
        footer_row = len(rows) + 2
        ws.cell(row=footer_row, column=1).value = "GRAND TOTAL"
        ws.cell(row=footer_row, column=1).font = Font(bold=True)
        ws.cell(row=footer_row, column=len(SERVICE_TYPES) + 2).value = total_quantity
        ws.cell(row=footer_row, column=len(SERVICE_TYPES) + 2).font = Font(bold=True)

        # Adjust column widths
        ws.column_dimensions["A"].width = 20
        for col_idx in range(2, len(SERVICE_TYPES) + 3):
            ws.column_dimensions[chr(64 + col_idx)].width = 15

        # Write to bytes
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"output_totals_{timestamp}.xlsx"
        return send_file(
            output,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            as_attachment=True,
            download_name=filename,
        )

    @app.route("/city-records")
    def city_records():
        start_date, end_date = _resolve_date_filter()
        rows = city_service_rows(
            {"start_date": start_date, "end_date": end_date},
            city_municipalities=CITY_MUNICIPALITIES,
        )
        return render_template(
            "city_records.html",
            month=start_date[:7],
            start_date=start_date,
            end_date=end_date,
            rows=rows,
            service_types=SERVICE_TYPES,
        )

    @app.get("/reports/city-records")
    def city_records_report():
        start_date, end_date = _resolve_date_filter()
        rows = city_service_rows(
            {"start_date": start_date, "end_date": end_date},
            city_municipalities=CITY_MUNICIPALITIES,
        )
        pdf_path = generate_city_service_summary_report(rows, SERVICE_TYPES, start_date, end_date)
        return send_file(pdf_path, mimetype="application/pdf")

    @app.route("/schedules")
    def schedules():
        filters = {
            "start_date": (request.args.get("start_date") or "").strip(),
            "end_date": (request.args.get("end_date") or "").strip(),
            "status": (request.args.get("status") or "").strip(),
            "employee_id": (request.args.get("employee_id") or "").strip(),
        }
        filters = {key: value for key, value in filters.items() if value}
        rows = list_schedules(filters)
        calendar_events = [
            {
                "title": f"{row['city_municipality']} - {(row['status'] or 'Pending')}",
                "start": row["schedule_date"],
                "color": "#16a34a" if (row["status"] or "").lower() == "approved" else "#f59e0b",
                "url": url_for("schedule_form", schedule_id=row["id"]),
            }
            for row in rows
        ]
        return render_template(
            "schedules.html",
            schedules=rows,
            filters=filters,
            calendar_events=calendar_events,
            employees=list_employees(),
        )

    @app.get("/reports/schedules")
    def schedules_report():
        filters = {
            "start_date": (request.args.get("start_date") or "").strip(),
            "end_date": (request.args.get("end_date") or "").strip(),
            "status": (request.args.get("status") or "").strip(),
            "employee_id": (request.args.get("employee_id") or "").strip(),
        }
        scoped = {k: v for k, v in filters.items() if v}
        rows = list_schedules(scoped)
        pdf_path = generate_schedule_report(
            rows,
            start_date=filters["start_date"],
            end_date=filters["end_date"],
            status=filters["status"],
        )
        return send_file(pdf_path, mimetype="application/pdf")

    @app.get("/reports/locators-chart")
    def locator_chart_report():
        employee_id_raw = (request.args.get("employee_id") or "").strip()
        if not employee_id_raw:
            flash("Please select an employee before printing Locator's Chart.", "error")
            return redirect(url_for("schedules"))
        try:
            employee_id = int(employee_id_raw)
        except ValueError:
            flash("Invalid employee selected.", "error")
            return redirect(url_for("schedules"))

        employee = get_employee(employee_id)
        if not employee:
            flash("Employee not found.", "error")
            return redirect(url_for("schedules"))

        filters = {
            "start_date": (request.args.get("start_date") or "").strip(),
            "end_date": (request.args.get("end_date") or "").strip(),
            "status": (request.args.get("status") or "").strip(),
            "employee_id": employee_id_raw,
        }
        scoped = {k: v for k, v in filters.items() if v}
        rows = list_schedules(scoped)
        pdf_path = generate_locator_chart_report(
            dict(employee),
            rows,
            start_date=filters["start_date"],
            end_date=filters["end_date"],
        )
        return send_file(pdf_path, mimetype="application/pdf")

    @app.get("/reports/locators-chart-all")
    def locator_chart_all_report():
        filters = {
            "start_date": (request.args.get("start_date") or "").strip(),
            "end_date": (request.args.get("end_date") or "").strip(),
            "status": (request.args.get("status") or "").strip(),
        }
        scoped = {k: v for k, v in filters.items() if v}
        rows = list_schedules(scoped)
        by_employee: dict[int, list[dict]] = {}
        for row in rows:
            for emp_id in row.get("assigned_rko_ids", []):
                by_employee.setdefault(int(emp_id), []).append(row)
            for emp_id in row.get("assigned_ra_ids", []):
                by_employee.setdefault(int(emp_id), []).append(row)
        employees = [dict(emp) for emp in list_employees() if int(emp["id"]) in by_employee]
        pdf_path = generate_locator_chart_all_report(
            employees,
            by_employee,
            start_date=filters["start_date"],
            end_date=filters["end_date"],
        )
        return send_file(pdf_path, mimetype="application/pdf")

    @app.route("/schedules/new", methods=["GET", "POST"])
    @app.route("/schedules/<int:schedule_id>/edit", methods=["GET", "POST"])
    def schedule_form(schedule_id=None):
        schedule = get_schedule(schedule_id) if schedule_id else None
        if request.method == "POST":
            save_schedule(schedule_id, request.form)
            flash("Schedule saved successfully.", "success")
            return redirect(url_for("schedules"))
        existing_rows = list_schedules({})
        assigned_by_date: dict[str, dict[str, list[int]]] = {}
        for row in existing_rows:
            if schedule_id and row["id"] == schedule_id:
                continue
            date_key = (row["schedule_date"] or "").strip()
            if not date_key:
                continue
            bucket = assigned_by_date.setdefault(date_key, {"rko_ids": [], "ra_ids": []})
            for employee_id in row.get("assigned_rko_ids", []):
                bucket["rko_ids"].append(int(employee_id))
            for employee_id in row.get("assigned_ra_ids", []):
                bucket["ra_ids"].append(int(employee_id))
        rko_employees = list_schedule_employees("Registration Kit Operator")
        ra_employees = list_schedule_employees("Registration Assistant")
        return render_template(
            "schedule_form.html",
            schedule=schedule,
            rko_employees=rko_employees,
            ra_employees=ra_employees,
            assigned_by_date=assigned_by_date,
            rko_total=len(rko_employees),
            ra_total=len(ra_employees),
        )

    @app.post("/schedules/<int:schedule_id>/delete")
    def schedule_delete(schedule_id):
        delete_schedule(schedule_id)
        flash("Schedule deleted.", "success")
        return redirect(url_for("schedules"))

    @app.route("/signatories")
    def signatories():
        return render_template("signatories.html", signatories=list_signatories())

    @app.route("/signatories/new", methods=["GET", "POST"])
    @app.route("/signatories/<int:signatory_id>/edit", methods=["GET", "POST"])
    def signatory_form(signatory_id=None):
        signatory = get_signatory(signatory_id) if signatory_id else None
        if request.method == "POST":
            save_signatory(signatory_id, request.form)
            flash("Signatory saved successfully.", "success")
            return redirect(url_for("signatories"))
        return render_template("signatory_form.html", signatory=signatory)

    @app.post("/signatories/<int:signatory_id>/delete")
    def signatory_delete(signatory_id):
        delete_signatory(signatory_id)
        flash("Signatory deleted.", "success")
        return redirect(url_for("signatories"))

    # -----------------------------------------------------------------
    # Import Sources (Administrator only)
    #
    # An admin wires a Google Sheet (or Apps Script endpoint) into the
    # system here. From then on every user across every page/report just
    # sees the data it feeds into employee_outputs -- there's nothing
    # per-user to configure. Each saved link shows up with a checkbox: if
    # the admin unchecks it, its rows stop counting everywhere instantly
    # (they're removed from employee_outputs); re-checking it re-pulls the
    # sheet and it counts again.
    # -----------------------------------------------------------------

    @app.get("/imports")
    @roles_required(ROLE_ADMIN)
    def import_sources_page():
        return render_template("import_sources.html", sources=list_import_sources())

    @app.post("/imports/google-sheet")
    @roles_required(ROLE_ADMIN)
    def import_google_sheet():
        original_csv_url = request.form.get("csv_url", "").strip()
        label = (request.form.get("label") or "").strip()
        try:
            csv_url = _validate_google_sheet_url(original_csv_url)
        except ValueError as e:
            flash(str(e), "error")
            return redirect(url_for("import_sources_page"))
        try:
            inserted = import_from_csv_url(
                csv_url,
                original_url=original_csv_url,
                label=label,
                created_by_user_id=current_user_id(),
                created_by_name=session.get("full_name") or session.get("username"),
            )
            flash(f"Imported {inserted} new output rows from Google Sheet. Existing rows were kept without duplication.", "success")
        except RuntimeError as e:
            flash(f"Import failed: {str(e)}", "error")
        except Exception as e:
            flash(f"An unexpected error occurred during import: {str(e)}", "error")
        return redirect(url_for("import_sources_page"))

    @app.post("/imports/apps-script")
    @roles_required(ROLE_ADMIN)
    def import_apps_script():
        original_script_url = request.form.get("script_url", "").strip()
        label = (request.form.get("label") or "").strip()
        method = request.form.get("method", "GET").strip().upper()
        payload_text = request.form.get("payload", "").strip()
        try:
            _validate_apps_script_url(original_script_url)
        except ValueError as e:
            flash(str(e), "error")
            return redirect(url_for("import_sources_page"))
        payload = json.loads(payload_text) if payload_text else {}
        try:
            inserted = import_from_apps_script(
                original_script_url,
                method=method,
                payload=payload,
                original_url=original_script_url,
                label=label,
                created_by_user_id=current_user_id(),
                created_by_name=session.get("full_name") or session.get("username"),
            )
            flash(f"Imported {inserted} output rows from Apps Script.", "success")
        except RuntimeError as e:
            flash(f"Import failed: {str(e)}", "error")
        except Exception as e:
            flash(f"An unexpected error occurred during import: {str(e)}", "error")
        return redirect(url_for("import_sources_page"))

    @app.post("/imports/check-sources")
    @roles_required(ROLE_ADMIN)
    def check_import_sources():
        sources = [s for s in list_import_sources() if s["is_active"]]
        checked = 0
        failed = 0
        for source in sources:
            try:
                reimport_source(source)
                checked += 1
            except Exception:
                failed += 1
        summary = f"Checked {len(sources)} active saved link(s). {checked} successful, {failed} failed."
        flash(summary, "success" if failed == 0 else "warning")
        return redirect(url_for("import_sources_page"))

    @app.post("/imports/<int:source_id>/toggle")
    @roles_required(ROLE_ADMIN)
    def toggle_import_source(source_id):
        turn_on = request.form.get("is_active") == "1"
        source = set_import_source_active(source_id, turn_on)
        if not source:
            flash("That saved link no longer exists.", "error")
            return redirect(url_for("import_sources_page"))

        if not turn_on:
            removed = delete_outputs_by_source_ref(source["normalized_url"])
            flash(
                f"Unchecked \"{source['label'] or source['original_url']}\" -- {removed} row(s) no longer count anywhere in the system.",
                "success",
            )
        else:
            try:
                inserted = reimport_source(source)
                flash(
                    f"Re-checked \"{source['label'] or source['original_url']}\" -- {inserted} row(s) pulled back in and counting again.",
                    "success",
                )
            except Exception as e:
                flash(f"Re-enabled, but the re-check failed: {e}", "warning")
        return redirect(url_for("import_sources_page"))

    @app.post("/imports/<int:source_id>/delete")
    @roles_required(ROLE_ADMIN)
    def delete_import_source_route(source_id):
        source = get_import_source(source_id)
        if not source:
            flash("That saved link no longer exists.", "error")
            return redirect(url_for("import_sources_page"))
        removed = delete_outputs_by_source_ref(source["normalized_url"])
        delete_import_source(source_id)
        flash(
            f"Removed \"{source['label'] or source['original_url']}\" and {removed} row(s) it had contributed.",
            "success",
        )
        return redirect(url_for("import_sources_page"))

    @app.post("/imports/clear")
    @roles_required(ROLE_ADMIN)
    def clear_imported_data():
        source_type = request.form.get("source_type", "google_sheet").strip()
        deleted = delete_imported_outputs(source_type)
        if source_type == "google_sheet":
            label = "Google Sheet"
        elif source_type == "apps_script":
            label = "Apps Script"
        else:
            label = "imported"
        flash(f"Deleted {deleted} {label} imported rows from the database.", "success")
        return redirect(url_for("import_sources_page"))

    @app.post("/branding/logo")
    def upload_logo():
        uploaded = request.files.get("logo")
        if not uploaded or not uploaded.filename:
            flash("Please choose a logo file.", "error")
            return redirect(url_for("dashboard"))

        filename = secure_filename(uploaded.filename)
        ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
        if ext not in {"png", "jpg", "jpeg"}:
            flash("Logo must be a PNG or JPG image.", "error")
            return redirect(url_for("dashboard"))

        upload_dir = Path(app.config["UPLOAD_DIR"])
        for existing in upload_dir.glob("logo.*"):
            existing.unlink(missing_ok=True)
        saved_path = upload_dir / f"logo.{ext}"
        uploaded.save(saved_path)
        flash("Logo updated.", "success")
        return redirect(url_for("dashboard"))

    @app.post("/settings")
    def save_settings():
        for key in ("organization_name", "report_title", "trn_logsheet_url"):
            if key in request.form:
                set_setting(key, request.form.get(key, "").strip())
        flash("Report settings updated.", "success")
        return redirect(url_for("dashboard"))

    @app.get("/reports/employee/<int:employee_id>")
    def employee_report(employee_id):
        start_date, end_date = _resolve_date_filter()
        report_type = request.args.get("category") or "registration"
        pdf_path = generate_employee_report_by_date(employee_id, start_date, end_date, report_type)
        employee = get_employee(employee_id)
        return send_file(
            pdf_path,
            as_attachment=True,
            download_name=f"{start_date}_to_{end_date}_{employee['full_name']}_{report_type}.pdf",
        )

    @app.get("/reports/export-all")
    def export_all_reports():
        start_date, end_date = _resolve_date_filter()
        report_type = request.args.get("category") or "registration"
        zip_path = generate_bulk_reports_by_date(start_date, end_date, report_type)
        return send_file(zip_path, as_attachment=True, download_name=zip_path.name)

    @app.get("/reports/output-summary")
    def output_summary_report():
        filters, start_date, end_date, employee_id, employee_name = _output_report_filters_from_request()
        rows = output_per_person_rows(filters)
        prepared_position = None
        if employee_id:
            try:
                emp = get_employee(int(employee_id))
                prepared_position = emp["position"] if emp else None
            except (ValueError, TypeError):
                prepared_position = None
        pdf_path = generate_output_summary_report(
            rows,
            SERVICE_TYPES,
            start_date,
            end_date,
            employee_name,
            prepared_position,
        )
        # Allow previewing inline in the browser (default) or forcing download with ?download=1
        download = str(request.args.get("download", "")).lower() in ("1", "true", "yes")
        if download:
            resp = send_file(pdf_path, as_attachment=True, download_name=pdf_path.name)
        else:
            resp = send_file(pdf_path, mimetype="application/pdf")
        resp.headers["Cache-Control"] = "no-store, no-cache, must-revalidate, max-age=0"
        resp.headers["Pragma"] = "no-cache"
        return resp

    @app.get("/reports/output-summary/preview")
    def output_summary_preview():
        filters, start_date, end_date, employee_id, employee_name = _output_report_filters_from_request()
        rows = output_per_person_rows(filters)
        grand_totals = output_summary_grand_totals(rows, SERVICE_TYPES) if rows else None
        sigs = signatories_for_report("output")
        prepared = sigs.get("prepared_by")
        verified = sigs.get("verified_by")
        report_query = {}
        if employee_id:
            report_query["employee_id"] = employee_id
        if start_date:
            report_query["start_date"] = start_date
        if end_date:
            report_query["end_date"] = end_date
        return render_template(
            "output_summary_preview.html",
            rows=rows,
            grand_totals=grand_totals,
            service_types=SERVICE_TYPES,
            start_date=start_date,
            end_date=end_date,
            employee_id=employee_id,
            report_query=report_query,
            employee_name=employee_name,
            prepared_by_name=employee_name or ((prepared or {}).get("name") or ""),
            prepared_by_position=(prepared or {}).get("position") or "National ID Registration",
            verified_by=verified,
        )

    @app.get("/reports/dar")
    def dar_report():
        start_date, end_date = _resolve_output_report_date_range()
        employee_id = request.args.get("employee_id") or None
        employee_name = None
        if employee_id:
            try:
                emp = get_employee(int(employee_id))
                employee_name = emp["full_name"] if emp else None
            except (ValueError, TypeError):
                employee_name = None
        filters = {}
        if employee_id:
            filters["employee_id"] = employee_id
        if start_date:
            filters["start_date"] = start_date
        if end_date:
            filters["end_date"] = end_date
        rows = output_per_person_rows(filters)
        grand_totals = output_summary_grand_totals(rows, SERVICE_TYPES) if rows else None
        report_query = {}
        if employee_id:
            report_query["employee_id"] = employee_id
        if start_date:
            report_query["start_date"] = start_date
        if end_date:
            report_query["end_date"] = end_date
        sigs = signatories_for_report("output")
        return render_template(
            "dar_preview.html",
            rows=rows,
            employees=list_employees(),
            service_types=SERVICE_TYPES,
            start_date=start_date,
            end_date=end_date,
            employee_id=employee_id,
            employee_name=employee_name,
            grand_totals=grand_totals,
            report_query=report_query,
            prepared_by=sigs.get("prepared_by"),
            verified_by=sigs.get("verified_by"),
        )

    @app.get("/reports/dar/export-all-pdf")
    def dar_export_all_pdf():
        start_date, end_date = _resolve_output_report_date_range()
        zip_path = generate_all_employees_dar_pdf_zip(start_date, end_date)
        return send_file(
            zip_path,
            as_attachment=True,
            download_name=f"{start_date or 'all'}_to_{end_date or 'all'}_all_employees_DAR_pdf.zip",
        )

    @app.post("/reports/dar/download")
    def dar_download():
        start_date = (request.form.get("start_date") or "").strip()
        end_date = (request.form.get("end_date") or "").strip()
        employee_id = request.form.get("employee_id") or None
        city_municipality = (request.form.get("city_municipality") or "").strip() or None
        employee_name = None
        if employee_id:
            try:
                emp = get_employee(int(employee_id))
                employee_name = emp["full_name"] if emp else None
            except (ValueError, TypeError):
                employee_name = None
        filters = {}
        if employee_id:
            filters["employee_id"] = employee_id
        if start_date:
            filters["start_date"] = start_date
        if end_date:
            filters["end_date"] = end_date
        if city_municipality:
            filters["city_municipality"] = city_municipality
        rows = output_per_person_rows(filters)

        out = build_dar_workbook(
            rows,
            start_date=start_date,
            end_date=end_date,
            employee_id=employee_id,
            employee_name=employee_name,
        )
        filename = f"DAR_{(employee_id or 'all')}_{start_date or 'all'}_{end_date or 'all'}.xlsx"
        return send_file(
            out,
            as_attachment=True,
            download_name=filename,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    # -----------------------------------------------------------------
    # User accounts (Administrator only)
    # -----------------------------------------------------------------

    @app.get("/users")
    @roles_required(ROLE_ADMIN)
    def users_list():
        return render_template("users.html", users=list_users(), roles=ROLES)

    @app.route("/users/new", methods=["GET", "POST"])
    @app.route("/users/<int:user_id>/edit", methods=["GET", "POST"])
    @roles_required(ROLE_ADMIN)
    def user_form(user_id=None):
        user = get_user(user_id) if user_id else None
        if request.method == "POST":
            try:
                if user_id:
                    update_user(
                        user_id,
                        full_name=request.form.get("full_name"),
                        role=request.form.get("role"),
                        active=bool(request.form.get("active")),
                        new_password=(request.form.get("password") or "").strip() or None,
                    )
                    flash("User updated.", "success")
                else:
                    create_user(
                        username=request.form.get("username"),
                        password=request.form.get("password"),
                        full_name=request.form.get("full_name"),
                        role=request.form.get("role"),
                    )
                    flash("User created.", "success")
                return redirect(url_for("users_list"))
            except (ValueError, sqlite3.IntegrityError) as exc:
                flash(f"Could not save user: {exc}", "error")
        return render_template("user_form.html", user=user, roles=ROLES)

    # Note: user accounts are never deleted once registered. To remove
    # someone's access, edit their account and uncheck "Active" instead --
    # this keeps their history (messages, reported concerns, outputs, etc.)
    # intact while disabling their login.

    # -----------------------------------------------------------------
    # 1-on-1 messenger -- messages auto-delete 10 hours after being sent
    # -----------------------------------------------------------------

    @app.get("/messages")
    def messages_home():
        uid = current_user_id()
        contacts = list_active_users(exclude_user_id=uid)
        unread = unread_dm_counts_by_sender(uid)
        return render_template("messages.html", contacts=contacts, unread=unread, active_contact=None, thread=[])

    @app.route("/messages/<int:user_id>", methods=["GET", "POST"])
    def messages_thread(user_id):
        uid = current_user_id()
        if user_id == uid:
            abort(404)
        partner = get_user(user_id)
        if not partner or not partner["active"]:
            abort(404)
        if request.method == "POST":
            send_direct_message(uid, user_id, request.form.get("body"))
            if request.headers.get("X-Requested-With") == "fetch":
                return ("", 204)
            return redirect(url_for("messages_thread", user_id=user_id))
        contacts = list_active_users(exclude_user_id=uid)
        unread = unread_dm_counts_by_sender(uid)
        return render_template(
            "messages.html",
            contacts=contacts,
            unread=unread,
            active_contact=partner,
            thread=list_conversation(uid, user_id),
        )

    @app.get("/messages/<int:user_id>/poll")
    def messages_poll(user_id):
        from flask import jsonify

        uid = current_user_id()
        thread = list_conversation(uid, user_id)
        return jsonify(
            [
                {
                    "sender_id": row["sender_id"],
                    "sender_name": row["sender_name"],
                    "body": row["body"],
                    "created_at": row["created_at"],
                    "mine": row["sender_id"] == uid,
                }
                for row in thread
            ]
        )


def _current_month():
    from datetime import datetime

    return datetime.now().strftime("%Y-%m")


def _resolve_output_report_date_range() -> tuple[str, str]:
    """Start/end from query string only. Empty means no date bound (same as Employee Output filters)."""
    start_date = (request.args.get("start_date") or "").strip()
    end_date = (request.args.get("end_date") or "").strip()
    return start_date, end_date


def _output_report_filters_from_request():
    """Build filters for ``output_per_person_rows`` from the report URL (preview / PDF)."""
    start_date, end_date = _resolve_output_report_date_range()
    employee_raw = (request.args.get("employee_id") or "").strip()
    filters: dict = {}
    if start_date:
        filters["start_date"] = start_date
    if end_date:
        filters["end_date"] = end_date
    employee_id = None
    employee_name = None
    if employee_raw:
        filters["employee_id"] = employee_raw
        try:
            emp = get_employee(int(employee_raw))
            employee_name = emp["full_name"] if emp else None
        except (ValueError, TypeError):
            employee_name = None
        employee_id = employee_raw
    return filters, start_date, end_date, employee_id, employee_name


def _resolve_date_filter():
    from datetime import datetime

    start_date = request.args.get("start_date") or request.form.get("start_date") or ""
    end_date = request.args.get("end_date") or request.form.get("end_date") or ""
    if start_date and end_date:
        return start_date, end_date

    today = datetime.now().strftime("%Y-%m-%d")
    return start_date or today, end_date or start_date or today


def _normalize_google_sheet_url(url: str) -> str:
    if not url:
        return ""
    if "/gviz/tq" in url or "export?format=csv" in url:
        return url
    if "/edit" in url and "/spreadsheets/d/" in url:
        sheet_id = url.split("/spreadsheets/d/")[1].split("/")[0]
        return f"https://docs.google.com/spreadsheets/d/{sheet_id}/gviz/tq?tqx=out:csv&sheet=Form%20Responses%201"
    return url


def _validate_google_sheet_url(original_url: str) -> str:
    """Normalize a Google Sheet link and make sure it's actually a Google
    Sheets URL before the server is told to fetch it -- only the assigned
    admin can wire a link in at all (route is admin-only), and this keeps
    the system from being pointed at an arbitrary third-party URL."""
    if not original_url:
        raise ValueError("Google Sheet CSV export URL is required.")
    from urllib.parse import urlparse

    parsed = urlparse(original_url)
    if parsed.scheme != "https" or parsed.netloc.lower() != "docs.google.com":
        raise ValueError("Only docs.google.com Google Sheets links are allowed.")
    csv_url = _normalize_google_sheet_url(original_url)
    if not csv_url:
        raise ValueError("Could not recognize that as a Google Sheet link.")
    return csv_url


def _validate_apps_script_url(url: str) -> None:
    if not url:
        raise ValueError("Apps Script URL is required.")
    from urllib.parse import urlparse

    parsed = urlparse(url)
    if parsed.scheme != "https" or parsed.netloc.lower() != "script.google.com":
        raise ValueError("Only script.google.com Apps Script links are allowed.")
